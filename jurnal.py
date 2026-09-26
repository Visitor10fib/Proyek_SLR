import os
import litstudy
import pandas as pd
from docx import Document

# Tentukan direktori kerja dan file korpus BibTeX Anda
folder = r"E:\MATA KULIAH SEMESTER 1 DEPARTEMEN GEOGRAFI UNIVERSITAS INDONESIA\Metodologi Penelitian Geografi - GeoMetPen\Proyek_SLR"
file_bib = os.path.join(folder, "Tes ref.bib")
output_docx = os.path.join(folder, "Hasil_Analisis_Litstudy_Final.docx")

print("==================================================")
print(" MENJALANKAN WORKFLOW LITSTUDY & EKSPOR LAPORAN ")
print("==================================================")

if os.path.exists(file_bib):
    # 1. Memuat file korpus format BibTeX (.bib)
    docs = litstudy.load_bibtex(file_bib)
    print(f"\n[1. MEMUAT DATA] Berhasil memuat {len(docs)} dokumen dari korpus literatur.")

    # 2. Statistik Dasar Publikasi per Tahun
    years = [d.publication_year for d in docs if hasattr(d, 'publication_year') and d.publication_year is not None]
    year_series = pd.Series(years).value_counts().sort_index()

    # 3. Membangun Korpus Teks untuk Natural Language Processing (NLP)
    print("\n[2. NLP & KORPUS] Membangun korpus teks dari judul dan abstrak...")
    corpus = litstudy.build_corpus(docs, ngram_threshold=0.6)
    print("[INFO] Korpus teks berhasil dibangun.")

    # 4. Melatih Model Topik menggunakan NMF
    num_topics = 5  
    print(f"\n[3. TOPIC MODELING] Melatih model NMF dengan {num_topics} topik utama...")
    topic_model = litstudy.train_nmf_model(corpus, num_topics)
    
    # 5. Membuat Dokumen Laporan Word (.docx)
    print(f"\n[4. EKSPOR LAPORAN] Menyimpan hasil analisis ke file Word...")
    doc = Document()
    doc.add_heading('Laporan Analisis Korpus Literatur (Litstudy Workflow)', 0)
    
    doc.add_paragraph(f'Total dokumen yang dianalisis dari korpus: {len(docs)} referensi.')
    
    # Tambahkan bagian Statistik Tahun
    doc.add_heading('Distribusi Publikasi per Tahun', level=1)
    table_year = doc.add_table(rows=1, cols=2)
    hdr_cells = table_year.rows[0].cells
    hdr_cells[0].text = 'Tahun'
    hdr_cells[1].text = 'Jumlah Publikasi'
    for yr, count in year_series.items():
        row_cells = table_year.add_row().cells
        row_cells[0].text = str(yr)
        row_cells[1].text = str(count)

    # Tambahkan bagian Topik NMF (menggunakan pemanggilan parameter standar litstudy)
    doc.add_heading('Hasil Ekstraksi Topik Riset Otomatis (NMF Topics)', level=1)
    for topic_idx in range(topic_model.num_topics):
        # Memanggil tanpa keyword argument 'top_n'
        tokens = topic_model.best_tokens_for_topic(topic_idx)
        # Ambil maksimal 10 token pertama jika hasilnya berupa list/tuple panjang
        tokens_list = list(tokens)[:10] if hasattr(tokens, '__iter__') else [str(tokens)]
        
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"Topik {topic_idx + 1}: ").bold = True
        p.add_run(", ".join(str(t) for t in tokens_list))

    # Simpan file dokumen
    doc.save(output_docx)
    print(f"\n[SELESAI] Laporan berhasil disimpan di: {output_docx}")
else:
    print(f"[GAGAL] File korpus BibTeX tidak ditemukan pada path: {file_bib}")